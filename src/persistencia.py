from openpyxl import Workbook, load_workbook
from pathlib import Path

from collections import deque

BASE_DIR = Path(__file__).resolve().parent.parent
DATA_DIR = BASE_DIR / "data"

Path.mkdir(DATA_DIR, exist_ok=True)


def criar_planilha():
    planilha = Workbook()
    sheet_original = planilha.active
    planilha.remove(sheet_original)

    return planilha

def criar_sheet(planilha, sheet_name):
    sheet = planilha.create_sheet(sheet_name)

    sheet.append([
        "Ciclo",
        "Média",
        "Mediana",
        "Menor",
        "Maior",
        "Eficiencia",
        "Completos"
    ])

    return sheet


def add_info_sheet(sheet, dados:deque):
    for dado in dados:
        sheet.append(list(dado.values()))


def salvar_planilha(nome_file, planilha):
    path = DATA_DIR / "resultados"
    Path.mkdir(path, exist_ok=True)
    planilha.save(path / nome_file)


def salvar_execucao(info_execucao:dict):
    path = Path(DATA_DIR / "execucoes.xlsx")

    if path.exists():
        planilha = load_workbook(path)
    else:
        planilha = Workbook()
        planilha.active.append(list(info_execucao.keys()))

    sheet = planilha.active
    
    sheet.append(list(info_execucao.values()))
    planilha.save(path)


def salvar_passos(dados):
    with open(DATA_DIR / "passos.txt", "w", encoding="utf-8") as arquivo:
        arquivo.write("".join(dados))

def salvar_qtable(dados):
    with open(DATA_DIR / "qtable.txt", "w", encoding="utf-8") as arquivo:
        arquivo.write("".join(dados))