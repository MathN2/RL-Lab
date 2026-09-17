from openpyxl import Workbook, load_workbook
from pathlib import Path

planilha = Workbook()
sheet_original = planilha.active
planilha.remove(sheet_original) # type: ignore

def criar_sheet(sheet_name):
    aba = planilha.create_sheet(sheet_name)

    aba.append([
        "Ciclo",
        "Média",
        "Mediana",
        "Menor",
        "Maior",
        "Eficiencia",
        "Completos"
    ])

    return planilha, aba


def salvar_execucao(epsilon_inicial, episodios, ciclos, completados, taxa_comp, eficiencia, passos):
    path = Path("data/execucoes.xlsx")

    if path.exists():
        planilha = load_workbook(path)
    else:
        planilha = Workbook()

    sheet = planilha.active
    sheet.append([
        "ε inicial",
        "episodios",
        "ciclos",
        "completados",
        "taxa de conclusão",
        "eficiencia media",
        "passos médios"
    ])

    sheet.append([
        epsilon_inicial,
        episodios,
        ciclos,
        completados,
        taxa_comp,
        eficiencia,
        passos
    ])
    planilha.save(path)


def salvar_resultados(planilha, sheet, ciclo, media, mediana, menor, maior, eficiencia, completos):
    sheet.append([ciclo, media, mediana, menor, maior, eficiencia, completos])

    planilha.save("data/resultados.xlsx")

def salvar_passos(dados):
    with open("data/passos.txt", "w", encoding="utf-8") as arquivo:
        arquivo.write("".join(dados))

def salvar_qtable(dados):
    with open("data/qtable.txt", "w", encoding="utf-8") as arquivo:
        arquivo.write("".join(dados))